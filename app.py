import re
import csv
from flask import Flask, request, jsonify, render_template, send_file
import os
from dotenv import load_dotenv
import uuid
import shutil
import requests
import logging
from functools import wraps
from werkzeug.utils import secure_filename
import fitz
from docx import Document
from datetime import datetime, timedelta
import threading
import time
import schedule
from parsers import parsers

load_dotenv()

app = Flask(__name__, template_folder='app/templates')

logging.basicConfig(level=logging.DEBUG)
logging.getLogger('pdfminer').setLevel(logging.WARNING)

def handle_errors(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        try:
            return func(*args, **kwargs)
        except Exception as e:
            logging.error(f"Error in {func.__name__}: {e}")
            raise
    return wrapper

class JobManager:
    def __init__(self):
        self.job_data = {
            "Job ID": str(uuid.uuid4()),
            "Temp Dir": "",
            "Documents": [],
            "Queries": [],
            "Responses": {},
            "Status": "Initialized",
            "Export Format": "",
            "Orientation": ""
        }

    @handle_errors
    def initialize_temp_dir(self):
        script_dir = os.path.dirname(os.path.abspath(__file__))
        temp_root_dir = os.path.join(script_dir, "temp")
        os.makedirs(temp_root_dir, exist_ok=True)
        temp_job_dir = os.path.join(temp_root_dir, f"{self.job_data['Job ID']}_temp")
        os.makedirs(temp_job_dir, exist_ok=True)
        self.job_data["Temp Dir"] = temp_job_dir
        logging.info(f"Temporary directory created: {temp_job_dir}")

    def update_status(self, new_status):
        self.job_data["Status"] = new_status
        logging.info(f"Job status updated to: {new_status}")

    @handle_errors
    def finalize_extraction(self):
        temp_dir = self.job_data.get("Temp Dir")
        logging.info(f"Attempting to delete temporary directory: {temp_dir}")
        if temp_dir and os.path.exists(temp_dir):
            try:
                shutil.rmtree(temp_dir)
                logging.info(f"Temporary directory {temp_dir} deleted successfully.")
            except Exception as e:
                logging.error(f"Error while deleting temporary directory {temp_dir}: {e}")
            finally:
                temp_root_dir = os.path.dirname(temp_dir)
                if not os.listdir(temp_root_dir):
                    try:
                        os.rmdir(temp_root_dir)
                        logging.info(f"Root directory {temp_root_dir} deleted successfully.")
                    except Exception as e:
                        logging.error(f"Error while deleting root directory {temp_root_dir}: {e}")
            self.update_status("Data Extraction Complete")
            logging.info("Temporary directory cleanup completed.")

    def get_job_data(self):
        return self.job_data

class DocumentHandler:
    def __init__(self, job_manager):
        self.job_manager = job_manager

    @handle_errors
    def process_documents(self):
        temp_dir = self.job_manager.job_data["Temp Dir"]
        for document in self.job_manager.job_data["Documents"]:
            path = document["Path"]
            alias = os.path.splitext(document["Alias"])[0]
            ext = document["Ext"]
            text = self.extract_text_from_file(path, ext)
            cleaned_text = self.clean_text(text)
            txt_path = os.path.join(temp_dir, alias + '.txt')
            with open(txt_path, 'w', encoding='utf-8') as txt_file:
                txt_file.write(cleaned_text)
            document["Path"] = txt_path
            logging.info(f"Processed document: {document}")
        self.job_manager.update_status("Documents Processed")
        logging.info("Documents processed successfully")

    @handle_errors
    def extract_text_from_file(self, path, ext):
        parse_methods = {
            '.pdf': self.parse_pdf,
            '.docx': self.parse_docx,
            '.txt': self.parse_txt
        }
        if f".{ext}" in parse_methods:
            return parse_methods[f".{ext}"](path)
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

    @handle_errors
    def parse_txt(self, path):
        with open(path, "r", encoding="utf-8") as file:
            return file.read()

    @handle_errors
    def parse_docx(self, path):
        doc = Document(path)
        return "\n".join([paragraph.text for paragraph in doc.paragraphs])

    @handle_errors
    def parse_pdf(self, file_path):
        text = ""
        document = fitz.open(file_path)
        for page_num in range(len(document)):
            page = document.load_page(page_num)
            text += page.get_text()
        return text

    def clean_text(self, text):
        text = re.sub(r'[^\x20-\x7E]+', ' ', text)
        return text

    @handle_errors
    def load_text(self, file_path):
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            logging.warning(f"UTF-8 decoding failed for {file_path}, trying ISO-8859-1 encoding")
            with open(file_path, 'r', encoding='ISO-8859-1') as file:
                return file.read()

class QueryManager:
    def __init__(self, job_manager):
        self.job_manager = job_manager

class DataExtractor:
    def __init__(self, job_manager, document_handler, query_manager, llm_interface):
        self.job_manager = job_manager
        self.document_handler = document_handler
        self.query_manager = query_manager
        self.llm_interface = llm_interface

    @handle_errors
    def initialize_job(self):
        self.job_manager.initialize_temp_dir()
        self.job_manager.update_status("Initialized")

    @handle_errors
    def process_documents(self):
        self.document_handler.process_documents()

    @handle_errors
    def process_queries_and_collect_responses(self):
        for document in self.job_manager.job_data["Documents"]:
            doc_text = self.document_handler.load_text(document["Path"])
            document_alias = document["Alias"]
            if document_alias not in self.job_manager.job_data["Responses"]:
                self.job_manager.job_data["Responses"][document_alias] = {}
            for query in self.job_manager.job_data["Queries"]:
                query_alias = query["Alias"]
                query_text = query["Text"]
                query_format = query["Format"]
                parser_entry = parsers.get(query_format)
                parser = parser_entry["parser"]
                json_required = parser_entry["json_required"]
                format_guidance_message = parser_entry["format_guidance"]

                response_format = None
                if json_required:
                    response_format = {"type": "json_object"}

                try:
                    response = self.llm_interface.ask(query_text, doc_text, parser, format_guidance_message, response_format)
                    if response is None:
                        response = "Information not available"
                    self.job_manager.job_data["Responses"][document_alias][query_alias] = response
                except ValueError as ve:
                    logging.error(f"ValueError in processing query '{query_alias}' for document '{document_alias}': {ve}")
                    self.job_manager.job_data["Responses"][document_alias][query_alias] = "Information not available"
                except Exception as e:
                    logging.error(f"Error in processing query '{query_alias}' for document '{document_alias}': {e}")
                    self.job_manager.job_data["Responses"][document_alias][query_alias] = "Information not available"
        self.job_manager.update_status("Queries Processed, Responses Collected")

    @handle_errors
    def run(self):
        self.initialize_job()
        self.process_documents()
        self.process_queries_and_collect_responses()
        logging.info("Responses collected.")
        return self.job_manager.get_job_data()

class OutputGenerator:
    def __init__(self, job_manager):
        self.job_manager = job_manager

    @handle_errors
    def ensure_output_directory_exists(self, output_path):
        output_dir = os.path.dirname(output_path)
        os.makedirs(output_dir, exist_ok=True)
        logging.info(f"Output directory ensured: {output_dir}")

    @handle_errors
    def generate_output(self):
        job_data = self.job_manager.get_job_data()
        job_id = job_data["Job ID"]
        output_format = job_data["Export Format"]
        temp_dir = job_data["Temp Dir"]
        output_path = os.path.join(temp_dir, f'{job_id}.{output_format}')
        
        self.ensure_output_directory_exists(output_path)
        
        if output_format == 'csv':
            self.generate_csv(output_path)
        elif output_format == 'docx':
            self.generate_docx(output_path)
        else:
            raise ValueError("Unrecognized output format specified. Use 'csv' or 'docx'.")

    @handle_errors
    def generate_csv(self, output_path):
        headers, rows = [], []
        job_data = self.job_manager.get_job_data()
        orientation = job_data["Orientation"]

        if orientation == "doc_row":
            headers = ["Documents"] + [q["Alias"] for q in job_data["Queries"]]
            for doc_alias, responses in job_data["Responses"].items():
                row = [doc_alias] + [responses.get(q["Alias"], "") for q in job_data["Queries"]]
                rows.append(row)
        elif orientation == "query_row":
            headers = ["Queries"] + [d["Alias"] for d in job_data["Documents"]]
            for query in job_data["Queries"]:
                row = [query["Alias"]]
                for doc in job_data["Documents"]:
                    response = job_data["Responses"].get(doc["Alias"], {}).get(query["Alias"], "")
                    row.append(response)
                rows.append(row)
        else:
            raise ValueError("Unrecognized orientation specified.")

        with open(output_path, 'w', newline='') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(headers)
            writer.writerows(rows)
        self.job_manager.update_status("CSV Output Generated")
        logging.info("CSV output generated successfully.")

    @handle_errors
    def generate_docx(self, output_path):
        job_data = self.job_manager.get_job_data()
        doc = Document()
        orientation = job_data["Orientation"]

        if orientation == "doc_row":
            for query in job_data["Queries"]:
                doc.add_heading(query["Alias"], level=2)
                for doc_alias, responses in job_data["Responses"].items():
                    response = responses.get(query["Alias"], "")
                    doc.add_paragraph(f"{doc_alias}: {response}")
        elif orientation == "query_row":
            for doc_info in job_data["Documents"]:
                doc_alias = doc_info["Alias"]
                doc.add_heading(doc_alias, level=2)
                for query in job_data["Queries"]:
                    response = job_data["Responses"].get(doc_alias, {}).get(query["Alias"], "")
                    doc.add_paragraph(f"{query['Alias']}: {response}")
        else:
            raise ValueError("Unrecognized orientation specified.")

        doc.save(output_path)
        self.job_manager.update_status("DOCX Output Generated")
        logging.info("DOCX output generated successfully.")

class LLMInterface:
    def __init__(self, api_key, model="gpt-4o", temperature=0.0):
        self.api_key = api_key
        self.url = "https://api.openai.com/v1/chat/completions"
        self.headers = {"Content-Type": "application/json", "Authorization": f"Bearer {self.api_key}"}
        self.model = model
        self.temperature = temperature

    @handle_errors
    def ask(self, query, context, parser, format_guidance_message=None, response_format=None):
        messages = [
            {"role": "system", "content": context},
            {"role": "user", "content": query}
        ]

        if format_guidance_message:
            messages.append({"role": "user", "content": format_guidance_message})

        data = {
            "model": self.model,
            "messages": messages,
            "temperature": self.temperature
        }

        if response_format:
            data["response_format"] = response_format

        logging.debug(f"OpenAI API request payload: {data}")

        response = self._make_request(data)
        
        logging.debug(f"OpenAI API response: {response}")

        if parser:
            try:
                if callable(parser):
                    parsed_response = parser(response)
                elif hasattr(parser, 'parse'):
                    parsed_response = parser.parse(response)
                else:
                    raise ValueError("Invalid parser provided.")

                if parsed_response is None or parsed_response == "":
                    raise ValueError("Parsed response is invalid or empty")

                return parsed_response
            except ValueError as ve:
                logging.error(f"ValueError in ask: {ve}")
                return "Not available"
            except Exception as e:
                logging.error(f"Error in ask: {e}")
                return "Not available"
        else:
            return response

    def _make_request(self, data):
        response = requests.post(self.url, headers=self.headers, json=data)
        response.raise_for_status()
        return response.json()['choices'][0]['message']['content']


def cleanup_old_files():
    temp_root_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'temp')
    now = datetime.now()
    age_threshold = timedelta(hours=24)

    if os.path.exists(temp_root_dir):
        for folder in os.listdir(temp_root_dir):
            folder_path = os.path.join(temp_root_dir, folder)
            if os.path.isdir(folder_path):
                folder_age = now - datetime.fromtimestamp(os.path.getctime(folder_path))
                if folder_age > age_threshold:
                    try:
                        shutil.rmtree(folder_path)
                        logging.info(f"Deleted old temp directory: {folder_path}")
                    except Exception as e:
                        logging.error(f"Error deleting {folder_path}: {e}")

def run_cleanup_job():
    schedule.every(1).hour.do(cleanup_old_files)

    while True:
        schedule.run_pending()
        time.sleep(1)

job_store = {}

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/begin_extraction', methods=['POST'])
def begin_extraction():
    job_manager = JobManager()
    job_manager.initialize_temp_dir()

    files = request.files.getlist('files')
    logging.info(f"Received {len(files)} files")
    for i, file in enumerate(files):
        filename = secure_filename(file.filename)
        path = os.path.join(job_manager.job_data["Temp Dir"], filename)
        file.save(path)
        alias = request.form.get(f'doc_alias_{i}')
        ext = request.form.get(f'doc_ext_{i}')
        logging.info(f"File {i}: filename={filename}, alias={alias}, ext={ext}")
        job_manager.job_data["Documents"].append({
            "Path": path,
            "Alias": alias,
            "Ext": ext
        })

    queries = []
    query_index = 0
    while f'query_text_{query_index}' in request.form:
        text = request.form.get(f'query_text_{query_index}')
        alias = request.form.get(f'query_alias_{query_index}')
        format = request.form.get(f'query_format_{query_index}')
        logging.info(f"Query {query_index}: text={text}, alias={alias}, format={format}")
        if text and alias and format:
            queries.append({
                "Text": text,
                "Alias": alias,
                "Format": format
            })
        query_index += 1
    job_manager.job_data["Queries"] = queries

    logging.info(f"Documents: {job_manager.job_data['Documents']}")
    logging.info(f"Queries: {job_manager.job_data['Queries']}")

    job_manager.job_data["Export Format"] = request.form['export_format']
    job_manager.job_data["Orientation"] = request.form['orientation']

    document_handler = DocumentHandler(job_manager)
    query_manager = QueryManager(job_manager)

    api_key = os.getenv('OPENAI_API_KEY')
    if not api_key:
        raise ValueError("OPENAI_API_KEY not set in .env file")
    llm_interface = LLMInterface(api_key)

    data_extractor = DataExtractor(job_manager, document_handler, query_manager, llm_interface)
    job_data = data_extractor.run()

    output_generator = OutputGenerator(job_manager)
    output_generator.generate_output()

    job_store[job_manager.job_data["Job ID"]] = job_manager

    return jsonify(job_data)

@app.route('/download_results/<job_id>', methods=['GET'])
def download_results(job_id):
    logging.info(f"Download requested for job ID: {job_id}")

    job_manager = job_store.get(job_id)
    if not job_manager:
        logging.error(f"Job ID not found: {job_id}")
        return "Results file not found", 404

    job_data = job_manager.get_job_data()
    output_format = job_data["Export Format"]
    temp_dir = job_data["Temp Dir"]
    results_path = os.path.join(temp_dir, f'{job_id}.{output_format}')

    logging.info(f"Looking for results file at path: {results_path}")

    if not os.path.exists(results_path):
        logging.error(f"Results file not found at path: {results_path}")
        return "Results file not found", 404

    response = send_file(results_path, as_attachment=True, download_name=f'{job_id}.{output_format}')

    def cleanup_temp_dir():
        logging.info(f"Cleaning up temporary directory: {temp_dir}")
        try:
            job_manager.finalize_extraction()
            logging.info(f"Temporary directory {temp_dir} deleted successfully.")
        except Exception as e:
            logging.error(f"Error while cleaning up temporary directory: {e}")
        finally:
            if job_id in job_store:
                del job_store[job_id]
                logging.info(f"Job {job_id} removed from job_store.")

    response.call_on_close(cleanup_temp_dir)

    return response

if __name__ == "__main__":
    cleanup_thread = threading.Thread(target=run_cleanup_job, daemon=True)
    cleanup_thread.start()
    app.run(debug=True)