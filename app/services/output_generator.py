import os
import csv
import logging
from docx import Document
from functools import wraps

def handle_errors(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        try:
            return func(*args, **kwargs)
        except Exception as e:
            logging.error(f"Error in {func.__name__}: {e}")
            raise
    return wrapper

class OutputGenerator:
    def __init__(self, job_manager):
        self.job_manager = job_manager

    @handle_errors
    def ensure_output_directory_exists(self, output_path):
        output_dir = os.path.dirname(output_path)
        os.makedirs(output_dir, exist_ok=True)
        logging.info(f"Output directory ensured: {output_dir}")

    @handle_errors
    def generate_output(self):
        job_data = self.job_manager.get_job_data()
        job_id = job_data["Job ID"]
        output_format = job_data["Export Format"]
        temp_dir = job_data["Temp Dir"]
        output_path = os.path.join(temp_dir, f'{job_id}.{output_format}')
        
        self.ensure_output_directory_exists(output_path)
        
        if output_format == 'csv':
            self.generate_csv(output_path)
        elif output_format == 'docx':
            self.generate_docx(output_path)
        else:
            raise ValueError("Unrecognized output format specified. Use 'csv' or 'docx'.")

    @handle_errors
    def generate_csv(self, output_path):
        headers, rows = [], []
        job_data = self.job_manager.get_job_data()
        orientation = job_data["Orientation"]

        if orientation == "doc_row":
            headers = ["Documents"] + [q["Alias"] for q in job_data["Queries"]]
            for doc_alias, responses in job_data["Responses"].items():
                row = [doc_alias] + [responses.get(q["Alias"], "") for q in job_data["Queries"]]
                rows.append(row)
        elif orientation == "query_row":
            headers = ["Queries"] + [d["Alias"] for d in job_data["Documents"]]
            for query in job_data["Queries"]:
                row = [query["Alias"]]
                for doc in job_data["Documents"]:
                    response = job_data["Responses"].get(doc["Alias"], {}).get(query["Alias"], "")
                    row.append(response)
                rows.append(row)
        else:
            raise ValueError("Unrecognized orientation specified.")

        with open(output_path, 'w', newline='') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(headers)
            writer.writerows(rows)
        self.job_manager.update_status("CSV Output Generated")
        logging.info("CSV output generated successfully.")

    @handle_errors
    def generate_docx(self, output_path):
        job_data = self.job_manager.get_job_data()
        doc = Document()
        orientation = job_data["Orientation"]

        if orientation == "doc_row":
            for query in job_data["Queries"]:
                doc.add_heading(query["Alias"], level=2)
                for doc_alias, responses in job_data["Responses"].items():
                    response = responses.get(query["Alias"], "")
                    doc.add_paragraph(f"{doc_alias}: {response}")
        elif orientation == "query_row":
            for doc_info in job_data["Documents"]:
                doc_alias = doc_info["Alias"]
                doc.add_heading(doc_alias, level=2)
                for query in job_data["Queries"]:
                    response = job_data["Responses"].get(doc_alias, {}).get(query["Alias"], "")
                    doc.add_paragraph(f"{query['Alias']}: {response}")
        else:
            raise ValueError("Unrecognized orientation specified.")

        doc.save(output_path)
        self.job_manager.update_status("DOCX Output Generated")
        logging.info("DOCX output generated successfully.")